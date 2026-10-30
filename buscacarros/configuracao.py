"""Leitura do config.ini e da lista de buscas (buscas.json, editada pelo painel)."""
import configparser
import json
import os
import shutil
import uuid
from datetime import datetime
from pathlib import Path

from .texto import UFS, inteiro, norm, para_uf, sim_nao, tipo_cambio, tipo_combustivel



def _raiz():
    """Pasta da instalação (onde ficam config.ini, buscas.json, dados, logs e resultados).
    Com o iniciador, o código fica em versoes/<versão>/buscacarros e a raiz vem de BUSCACARRO_RAIZ."""
    if os.environ.get('BUSCACARRO_RAIZ'):
        return Path(os.environ['BUSCACARRO_RAIZ']).resolve()
    pacote_pai = Path(__file__).resolve().parent.parent
    if pacote_pai.parent.name == 'versoes':
        return pacote_pai.parent.parent
    return pacote_pai


RAIZ = _raiz()
ARQ_BUSCAS = RAIZ / 'buscas.json'


class ErroConfig(Exception):
    pass


# --------------------------------------------------------------------------- config.ini
def garantir_config():
    """Na primeira instalação cria o config.ini a partir do exemplo (atualizações não o sobrescrevem)."""
    alvo = RAIZ / 'config.ini'
    exemplo = RAIZ / 'config.exemplo.ini'
    if not alvo.exists() and exemplo.exists():
        shutil.copyfile(exemplo, alvo)
    if alvo.exists():
        _remover_secao_email(alvo)
    return alvo


def _minutos(valor, padrao=120):
    n = inteiro(valor)
    return padrao if n is None else max(0, n)


def ler_config(caminho=None):
    cp = configparser.ConfigParser()
    caminho = Path(caminho or garantir_config())
    if not caminho.exists():
        raise ErroConfig(f'Arquivo {caminho.name} não encontrado.')
    try:
        cp.read(caminho, encoding='utf-8-sig')
    except UnicodeDecodeError:  # salvo como ANSI pelo Bloco de Notas
        cp = configparser.ConfigParser()
        cp.read(caminho, encoding='cp1252')

    def g(sec, chave, padrao=''):
        try:
            return cp.get(sec, chave).strip()
        except (configparser.NoSectionError, configparser.NoOptionError):
            return padrao

    def b(sec, chave, padrao=True):
        v = sim_nao(g(sec, chave, ''))
        return padrao if v is None else v

    return {
        'execucao': {
            'intervalo_dias': inteiro(g('execucao', 'intervalo_dias', '2')) or 2,
            'mostrar_navegador': b('execucao', 'mostrar_navegador', False),
            'reaproveitar_minutos': _minutos(g('execucao', 'reaproveitar_leitura_minutos', '120')),
        },
        'limites': {
            'paginas': inteiro(g('limites', 'max_paginas_por_site', '15')) or 15,
            'anuncios': inteiro(g('limites', 'max_anuncios_por_site', '300')) or 300,
        },
        'sites': {k: b('sites', k, True) for k in ('webmotors', 'olx', 'mercadolivre', 'seminovos', 'napista', 'icarros')},
        'fipe': {'token': g('fipe', 'token')},
        'alertas': {'abaixo_fipe_pct': -(abs(float(g('alertas', 'alertar_abaixo_da_fipe_pct', '25') or 25)))},
    }


def _ler_texto_config(caminho):
    b = caminho.read_bytes()
    for enc in ('utf-8-sig', 'cp1252'):
        try:
            return b.decode(enc)
        except UnicodeDecodeError:
            continue
    return b.decode('utf-8', errors='replace')


def _remover_secao_email(caminho):
    """O envio de e-mail saiu do programa (1.3.1): tira a seção [email] do config.ini, para não deixar
    guardada uma senha de app do Gmail que não é mais usada. O resto do arquivo fica como está."""
    try:
        texto = _ler_texto_config(caminho)
    except OSError:
        return
    linhas = texto.splitlines()
    ini = next((i for i, l in enumerate(linhas) if l.strip().lower() == '[email]'), None)
    if ini is None:
        return
    # o comentário do cabeçalho que mandava configurar o e-mail pelo painel vira o texto novo
    linhas = ['; O carro buscado se escolhe pelo painel; o resto pode ser ajustado aqui.'
              if l.strip().startswith(';') and 'E-mail do resultado' in l else l for l in linhas]
    fim = next((i for i in range(ini + 1, len(linhas)) if linhas[i].strip().startswith('[')), len(linhas))
    # leva junto os comentários logo acima da seção que falam de e-mail
    while ini > 0 and linhas[ini - 1].strip().startswith(';') and 'mail' in linhas[ini - 1].lower():
        ini -= 1
    novas = linhas[:ini] + linhas[fim:]
    nl = '\r\n' if '\r\n' in texto else '\n'
    try:
        caminho.write_text(nl.join(novas).rstrip() + nl, encoding='utf-8', newline='')
    except OSError:
        pass


# --------------------------------------------------------------------------- lista de buscas
CAMPOS = ('marca', 'modelo', 'versao', 'ano_de', 'ano_ate', 'km_max', 'uf', 'cidade', 'preco_max',
          'cambio', 'combustivel', 'blindado', 'vendedor', 'sites')

# sites que a pessoa pode escolher no painel (o iCarros fica só no config.ini: está sem busca de anúncios)
SITES_ESCOLHA = {'webmotors': 'Webmotors', 'olx': 'OLX', 'mercadolivre': 'Mercado Livre',
                 'seminovos': 'Seminovos BH', 'napista': 'NaPista'}


def sites_escolhidos(valor):
    """Lista de chaves de SITES_ESCOLHA, ou None quando são todos."""
    if valor is None or valor == '':
        return None
    if isinstance(valor, str):
        valor = [v.strip() for v in valor.split(',')]
    escolha = [k for k in SITES_ESCOLHA if k in set(valor)]
    if not escolha:
        raise ErroConfig('Escolha pelo menos um site para buscar.')
    return None if len(escolha) == len(SITES_ESCOLHA) else escolha


def busca_no_site(busca, chave):
    """Se o site entra na busca deste carro. Com uma escolha de sites, só eles (o iCarros, que só confere se a
    busca dele voltou, fica de fora); sem escolha, todos, conforme o config.ini."""
    return not busca.get('sites') or chave in busca['sites']


def tipo_vendedor(valor):
    """'particular', 'loja' ou None (qualquer)."""
    v = norm(valor)
    if v in ('particular', 'particulares', 'pf', 'pessoa fisica'):
        return 'particular'
    if v in ('loja', 'lojas', 'lojista', 'lojistas', 'revenda', 'revendas', 'pj'):
        return 'loja'
    return None



def normalizar_busca(dados):
    """Valida e padroniza uma busca vinda do painel. Levanta ErroConfig com mensagem para o usuário."""
    def txt(k):
        v = dados.get(k)
        v = str(v).strip() if v is not None else ''
        return v or None

    def num(k, nome):
        v = dados.get(k)
        if v in (None, ''):
            return None
        n = inteiro(v)
        if n is None or n <= 0:
            raise ErroConfig(f'{nome}: use só números.')
        return n

    marca, modelo = txt('marca'), txt('modelo')
    if not marca or not modelo:
        raise ErroConfig('Informe a marca e o modelo.')
    uf_txt = txt('uf')
    uf = para_uf(uf_txt) if uf_txt else None
    if uf_txt and not uf:
        raise ErroConfig(f'Estado "{uf_txt}" não reconhecido.')
    cidade = txt('cidade')
    if cidade and not uf:
        raise ErroConfig('Para buscar numa cidade, escolha também o estado.')
    b = {
        'marca': marca, 'modelo': modelo, 'versao': txt('versao'),
        'ano_de': num('ano_de', 'Ano de'), 'ano_ate': num('ano_ate', 'Ano até'),
        'km_max': num('km_max', 'Km máximo'), 'uf': uf, 'cidade': cidade,
        'preco_max': num('preco_max', 'Preço máximo'),
        'cambio': tipo_cambio(dados.get('cambio')) if norm(dados.get('cambio')) not in ('', 'qualquer') else None,
        'combustivel': tipo_combustivel(dados.get('combustivel')) if norm(dados.get('combustivel')) not in ('', 'qualquer') else None,
        'blindado': dados.get('blindado') if isinstance(dados.get('blindado'), bool) else sim_nao(dados.get('blindado')),
        'vendedor': tipo_vendedor(dados.get('vendedor')),
        'sites': sites_escolhidos(dados.get('sites')),
    }
    for k in ('ano_de', 'ano_ate'):
        if b[k] and not 1950 <= b[k] <= datetime.now().year + 2:
            raise ErroConfig('Use anos com 4 dígitos, ex.: 2015.')
    if b['ano_de'] and b['ano_ate'] and b['ano_de'] > b['ano_ate']:
        raise ErroConfig('O "ano de" precisa ser menor ou igual ao "ano até".')
    return b


def _completar(b):
    b = dict(b)
    b.setdefault('id', uuid.uuid4().hex[:8])
    for k in CAMPOS:
        b.setdefault(k, None)
    b['rotulo'] = rotulo(b)
    b['titulo'] = f"{b['marca']} {b['modelo']}" + (f" {b['versao']}" if b['versao'] else '')
    b['detalhes'] = detalhes(b)
    return b


def ler_lista():
    """Busca atual (a lista tem no máximo um carro). Na primeira vez, importa a planilha antiga, se houver."""
    if not ARQ_BUSCAS.exists():
        importadas = _importar_planilha()
        salvar_lista(importadas[-1:] if importadas else [])
    try:
        dados = json.loads(ARQ_BUSCAS.read_text(encoding='utf-8'))
    except Exception as e:
        raise ErroConfig(f'O arquivo buscas.json está corrompido ({e}).')
    return [_completar(b) for b in dados.get('buscas', [])][-1:]


def salvar_lista(buscas):
    limpas = []
    for b in buscas:
        x = {k: b.get(k) for k in CAMPOS}
        x['id'] = b.get('id') or uuid.uuid4().hex[:8]
        x['criada_em'] = b.get('criada_em') or datetime.now().isoformat(timespec='minutes')
        limpas.append(x)
    tmp = ARQ_BUSCAS.with_suffix('.tmp')
    tmp.write_text(json.dumps({'buscas': limpas}, ensure_ascii=False, indent=1), encoding='utf-8')
    tmp.replace(ARQ_BUSCAS)


def definir_busca(dados):
    """Salva o carro buscado. Substitui o anterior (a busca é de um carro por vez).
    Retorna (nova, anterior_ou_None)."""
    nova = normalizar_busca(dados)
    atual = ler_lista()
    chave = lambda b: tuple(norm(b.get(k)) for k in CAMPOS)  # noqa: E731
    if atual and chave(atual[0]) == chave(nova):
        raise ErroConfig('Esse já é o carro buscado.')
    nova['id'] = uuid.uuid4().hex[:8]
    salvar_lista([nova])
    return _completar(nova), (atual[0] if atual else None)


def excluir_busca(id_busca):
    lista = ler_lista()
    resto = [b for b in lista if b['id'] != id_busca]
    if len(resto) == len(lista):
        raise ErroConfig('Esse carro não está mais na busca.')
    salvar_lista(resto)


def ler_buscas():
    """Buscas prontas para a coleta."""
    return ler_lista()


def _importar_planilha():
    """Migra a buscas.xlsx da versão anterior (uma vez) e a renomeia para não confundir."""
    xlsx = RAIZ / 'buscas.xlsx'
    if not xlsx.exists():
        return None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx, data_only=True, read_only=True)
        ws = wb['Buscas'] if 'Buscas' in wb.sheetnames else wb.worksheets[0]
        linhas = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception:
        return None
    colunas = {'ativo': 'ativo', 'marca': 'marca', 'modelo': 'modelo', 'versao': 'versao', 'ano de': 'ano_de',
               'ano ate': 'ano_ate', 'km maximo': 'km_max', 'estado (uf)': 'uf', 'uf': 'uf', 'estado': 'uf',
               'cidade': 'cidade', 'preco maximo': 'preco_max', 'cambio': 'cambio', 'combustivel': 'combustivel',
               'blindado': 'blindado'}
    mapa = {colunas[norm(h)]: i for i, h in enumerate(linhas[0] if linhas else []) if h and norm(h) in colunas}
    buscas = []
    for row in linhas[1:]:
        d = {k: (row[i] if i < len(row) else None) for k, i in mapa.items()}
        if sim_nao(d.get('ativo')) is False:
            continue
        try:
            buscas.append(normalizar_busca(d))
        except ErroConfig:
            continue
    try:
        xlsx.rename(RAIZ / 'buscas_planilha_antiga.xlsx')
    except OSError:
        pass
    return buscas


# --------------------------------------------------------------------------- textos
def detalhes(b):
    partes = []
    if b['ano_de'] or b['ano_ate']:
        if b['ano_de'] and b['ano_ate']:
            partes.append(f"{b['ano_de']} a {b['ano_ate']}")
        elif b['ano_de']:
            partes.append(f"a partir de {b['ano_de']}")
        else:
            partes.append(f"até {b['ano_ate']}")
    if b['cambio']:
        partes.append('automático' if b['cambio'] == 'automatico' else 'manual')
    if b['combustivel']:
        partes.append({'flex': 'flex', 'gasolina': 'gasolina', 'diesel': 'diesel', 'hibrido': 'híbrido',
                       'eletrico': 'elétrico', 'alcool': 'álcool'}.get(b['combustivel'], b['combustivel']))
    if b['blindado'] is True:
        partes.append('só blindados')
    elif b['blindado'] is False:
        partes.append('sem blindagem')
    if b.get('vendedor') == 'particular':
        partes.append('só particulares')
    elif b.get('vendedor') == 'loja':
        partes.append('só lojas')
    if b['km_max']:
        partes.append(f"até {b['km_max']:,} km".replace(',', '.'))
    if b['preco_max']:
        partes.append(f"até R$ {b['preco_max']:,}".replace(',', '.'))
    local = (f"{b['cidade']}/{b['uf']}" if b['cidade'] else (UFS.get(b['uf']) if b['uf'] else 'Brasil inteiro'))
    partes.append(local)
    if b.get('sites'):
        partes.append('só em ' + ', '.join(SITES_ESCOLHA[k] for k in b['sites'] if k in SITES_ESCOLHA))
    return ', '.join(partes)


def rotulo(b):
    return f"{b['marca']} {b['modelo']}" + (f" {b['versao']}" if b['versao'] else '') + ', ' + detalhes(b)
